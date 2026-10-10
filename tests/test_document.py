"""Tests for document processing."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from src.document import extract_text_from_file, chunk_text, iter_supported_files


def test_chunk_text_basic():
    text = "A" * 1200
    chunks = chunk_text(text, chunk_size=500, overlap=50)
    assert len(chunks) >= 2
    assert all(len(c) <= 550 for c in chunks)


def test_chunk_text_short():
    assert chunk_text("hello") == ["hello"]
    assert chunk_text("") == []
    assert chunk_text("   ") == []


def test_extract_txt(tmp_path: Path):
    f = tmp_path / "note.txt"
    f.write_text("Industrial note about reactor safety.", encoding="utf-8")
    text = extract_text_from_file(f)
    assert "reactor" in text.lower()


def test_extract_unsupported(tmp_path: Path):
    f = tmp_path / "data.bin"
    f.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text_from_file(f)


def test_extract_missing():
    with pytest.raises(FileNotFoundError):
        extract_text_from_file(Path("/nonexistent/file.txt"))


def test_extract_image_registers_metadata_or_optional_ocr(tmp_path: Path):
    from PIL import Image

    path = tmp_path / "tiny.png"
    Image.new("RGB", (2, 2), color=(0, 0, 0)).save(path)
    text = extract_text_from_file(path)
    assert "Image file: tiny.png" in text
    assert "2x2 px" in text


@pytest.mark.parametrize(
    ("chunk_size", "overlap"),
    [(0, 0), (-1, 0), (10, -1), (10, 10), (10, 11)],
)
def test_chunk_text_rejects_invalid_window(chunk_size: int, overlap: int):
    with pytest.raises(ValueError):
        chunk_text("text that would otherwise be chunked", chunk_size=chunk_size, overlap=overlap)


def test_extract_docx_includes_table_cells_in_document_order(tmp_path: Path):
    path = tmp_path / "manual.docx"
    doc = Document()
    doc.add_paragraph("Operating limits")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Temperature"
    table.cell(0, 1).text = "450 C"
    table.cell(1, 0).text = "Pressure"
    table.cell(1, 1).text = "8–12 bar"
    doc.add_paragraph("End of manual")
    doc.save(path)

    extracted = extract_text_from_file(path)

    assert "Temperature | 450 C" in extracted
    assert "Pressure | 8–12 bar" in extracted
    assert extracted.index("Operating limits") < extracted.index("Temperature | 450 C")
    assert extracted.index("Pressure | 8–12 bar") < extracted.index("End of manual")


def test_iter_supported_files_skips_symlinks_outside_directory(tmp_path: Path):
    root = tmp_path / "ingest-root"
    root.mkdir()
    outside = tmp_path / "private.txt"
    outside.write_text("private text", encoding="utf-8")
    link = root / "linked.txt"
    try:
        link.symlink_to(outside)
    except (OSError, NotImplementedError):
        pytest.skip("Symlinks are not available in this environment.")

    assert list(iter_supported_files(root)) == []


def test_pdf_without_extractable_text_returns_empty_string(tmp_path: Path):
    from pypdf import PdfWriter

    path = tmp_path / "scanned-without-text.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)

    assert extract_text_from_file(path) == ""
